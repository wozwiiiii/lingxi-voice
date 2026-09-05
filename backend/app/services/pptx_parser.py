"""PPTX text extraction using python-pptx. Only supports .pptx format."""

import logging
import zipfile
from dataclasses import dataclass

from pptx import Presentation

logger = logging.getLogger(__name__)


@dataclass
class SlideText:
    index: int  # 1-based
    text: str
    notes: str = ""
    image_alt_texts: list[str] | None = None


@dataclass
class PPTXParseResult:
    title: str
    slides: list[SlideText]
    slides_count: int


MAX_UNZIPPED_SIZE = 100 * 1024 * 1024  # 100MB defense against ZIP bombs


def parse_pptx(filepath: str, original_filename: str = "") -> PPTXParseResult:
    """Extract text from each slide in a .pptx file.

    Raises ValueError for unparseable files.
    """
    # ZIP bomb defense: check total uncompressed size
    try:
        with zipfile.ZipFile(filepath) as zf:
            total_size = sum(info.file_size for info in zf.infolist())
            if total_size > MAX_UNZIPPED_SIZE:
                raise ValueError(
                    f"PPTX 解压后过大({total_size / 1024 / 1024:.0f}MB > "
                    f"{MAX_UNZIPPED_SIZE / 1024 / 1024:.0f}MB)"
                )
    except FileNotFoundError:
        raise ValueError(f"文件不存在: {filepath}")
    except zipfile.BadZipFile:
        raise ValueError("文件不是有效的 ZIP 格式(PowerPoint 文件本质是 ZIP)")

    try:
        prs = Presentation(filepath)
    except Exception as e:
        raise ValueError(f"无法解析 PPTX 文件: {e}")

    title = original_filename.rsplit(".", 1)[0] if original_filename else "未命名课件"
    slides: list[SlideText] = []

    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        notes_text = ""
        alt_texts: list[str] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                frame_text = shape.text_frame.text.strip()
                if frame_text:
                    texts.append(frame_text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_texts = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]
                    if row_texts:
                        texts.append(" | ".join(row_texts))

        try:
            notes_slide = slide.notes_slide
            notes_text = notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        combined_text = "\n".join(texts).strip()
        slides.append(SlideText(
            index=i,
            text=combined_text,
            notes=notes_text,
            image_alt_texts=alt_texts if alt_texts else None,
        ))

    return PPTXParseResult(
        title=title,
        slides=slides,
        slides_count=len(slides),
    )
