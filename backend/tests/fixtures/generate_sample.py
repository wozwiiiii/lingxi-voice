"""Generate a minimal 2-slide .pptx for testing."""

from pptx import Presentation
from pptx.util import Inches

prs = Presentation()

slide1 = prs.slides.add_slide(prs.slide_layouts[1])
slide1.shapes.title.text = "第一页标题"
slide1.placeholders[1].text = "这是第一页的内容文本"

slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "第二页标题"
slide2.placeholders[1].text = "这是第二页的内容文本"

prs.save("sample.pptx")
print("Generated: sample.pptx")
